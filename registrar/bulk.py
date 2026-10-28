"""Reading a list of clergy names for the bulk account import, and checking
each one before anything is created.

The registrar can paste names (one per line — which is also what pasting a
single spreadsheet column gives you) or upload a .csv/.xlsx whose first column
is the full name. All three end up as the same list of strings and go through
the same checks, so the preview table looks identical whichever they used.
"""

import csv
import io
import os
import re

from .models import ClergyRecord

MAX_NAMES = 100
MAX_UPLOAD_BYTES = 1024 * 1024  # 1 MB — a 100-name sheet is a few KB
NAME_MAX_LENGTH = ClergyRecord._meta.get_field('name').max_length

# A header cell in the first row is skipped rather than imported as a person.
_HEADER_LABELS = {'name', 'full name', 'fullname', 'clergy name', 'clergy'}


class BulkInputError(Exception):
    """The input couldn't be read at all; the message is shown to the registrar."""


def _clean(value):
    return re.sub(r'\s+', ' ', str(value)).strip() if value is not None else ''


def _drop_header(names):
    if names and names[0].lower() in _HEADER_LABELS:
        return names[1:]
    return names


def _names_from_text(text):
    return _drop_header([n for n in (_clean(line) for line in text.splitlines()) if n])


def _names_from_csv(raw):
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise BulkInputError('That CSV is not UTF-8 text. Re-save it as "CSV UTF-8" and try again.')
    rows = csv.reader(io.StringIO(text))
    return _drop_header([n for n in (_clean(row[0]) if row else '' for row in rows) if n])


def _names_from_xlsx(raw):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError):
        raise BulkInputError('That file could not be read as an Excel workbook.')
    try:
        sheet = workbook.worksheets[0]
        names = []
        for row in sheet.iter_rows(min_row=1, max_col=1, values_only=True):
            name = _clean(row[0]) if row else ''
            if name:
                names.append(name)
                if len(names) > MAX_NAMES + 1:  # stop reading a runaway sheet early
                    break
    finally:
        workbook.close()
    return _drop_header(names)


def read_names(text='', upload=None):
    """Returns the list of cleaned, non-blank names from the pasted text or
    the uploaded file. An upload wins if both are given."""
    if upload is not None:
        if upload.size > MAX_UPLOAD_BYTES:
            raise BulkInputError('That file is too large. A list of names should be well under 1 MB.')
        raw = upload.read()
        extension = os.path.splitext(upload.name)[1].lower()
        if extension == '.xlsx':
            if not raw.startswith(b'PK\x03\x04'):  # .xlsx is a zip container
                raise BulkInputError('That file does not look like a real .xlsx workbook.')
            names = _names_from_xlsx(raw)
        elif extension == '.csv':
            names = _names_from_csv(raw)
        else:
            raise BulkInputError('Upload a .xlsx or .csv file, or paste the names instead.')
    else:
        names = _names_from_text(text or '')

    if not names:
        raise BulkInputError('No names found. Put one full name per line.')
    if len(names) > MAX_NAMES:
        raise BulkInputError(f'Too many names ({len(names)}). Import at most {MAX_NAMES} at a time.')
    return names


def build_preview(names):
    """One dict per name: {'name', 'errors', 'warnings'}. An error blocks that
    row from being created; a warning is shown but the registrar decides —
    two clergy can genuinely share a name."""
    existing = {n.lower() for n in ClergyRecord.objects.values_list('name', flat=True)}
    seen = set()
    rows = []
    for name in names:
        errors, warnings = [], []
        if len(name) > NAME_MAX_LENGTH:
            errors.append(f'Longer than {NAME_MAX_LENGTH} characters.')
        elif not re.search(r'[^\W\d_]', name):  # needs at least one letter
            errors.append('Does not look like a name.')
        else:
            key = name.lower()
            if key in existing:
                warnings.append('A record with this name already exists.')
            if key in seen:
                warnings.append('Appears more than once in this list.')
            seen.add(key)
        rows.append({'name': name, 'errors': errors, 'warnings': warnings})
    return rows


def valid_names(rows):
    return [r['name'] for r in rows if not r['errors']]
